# New Semester Observing Schedule
# Step 11: Staff Schedule Uploader
# This script converts an Excel file to CSV format and optionally uploads it to the staff schedule site.

import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog
import datetime
import importlib
import html
import re
import shutil
import ssl
import subprocess
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.poolmanager import PoolManager

# Allow importing shared utilities from project root when running from s11_staff_scheds/.
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from common.obs_sem_utils import (
    _normalize_month_range,
    _normalize_semester_token,
    _normalize_staff_type,
    build_output_filename,
    ensure_data_dir,
    copy_source_to_data_dir,
    store_last_csv_filename,
    load_live_config,
)
from common.db_tunnel_utils import (
    collect_mysql_overview,
    format_type_summary,
    run_with_mysql_connection,
    run_mysql_preflight,
    apply_insert_statements_to_database,
    write_insert_results_to_file,
    extract_insert_statements,
    save_insert_statements,
)


LIVE_CONFIG = load_live_config()
SSH_PASSWORD_CACHE = {}


def get_config(section, key, default=None):
    """Get a configuration value from LIVE_CONFIG."""
    return LIVE_CONFIG.get(section, {}).get(key, default)


def _is_debug_enabled():
    """Return True when debug logging is enabled via config."""
    value = get_config("NEW_OBS_SEM", "DEBUG", False)
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "on"}
    return bool(value)


def _debug_log(message):
    """Emit debug message only when DEBUG is enabled."""
    if _is_debug_enabled():
        print(f"[DEBUG] {message}")


def get_ssh_password(settings):
    """Get SSH password from cache or prompt user for it."""
    cache_key = (settings.get("ssh_host"), settings.get("ssh_user"), settings.get("ssh_port"))
    cached = SSH_PASSWORD_CACHE.get(cache_key)
    if cached:
        return cached

    parent = getattr(tk, "_default_root", None)
    display_host = settings.get("ssh_hostname") or settings.get("ssh_host")
    prompt = (
        f"Enter SSH password for {settings.get('ssh_user')}@{display_host}"
        #f" (port {settings.get('ssh_port')})"
    )
    password = simpledialog.askstring("SSH Password Required", prompt, show="*", parent=parent)
    if not password:
        raise RuntimeError("SSH password entry was cancelled.")

    SSH_PASSWORD_CACHE[cache_key] = password
    return password


class LegacyTLSAdapter(HTTPAdapter):
    def init_poolmanager(self, connections, maxsize, block=False, **pool_kwargs):
        context = ssl.create_default_context()
        context.set_ciphers("DEFAULT@SECLEVEL=1")
        self.poolmanager = PoolManager(
            num_pools=connections,
            maxsize=maxsize,
            block=block,
            ssl_context=context,
            **pool_kwargs,
        )


def _set_dialog_geometry(
    dialog,
    width,
    height,
    min_width=420,
    min_height=220,
    max_width=None,
    max_height=None,
):
    """Clamp dialog size to screen and center it so controls stay visible."""
    screen_width = dialog.winfo_screenwidth()
    screen_height = dialog.winfo_screenheight()

    screen_max_width = max(min_width, screen_width - 80)
    screen_max_height = max(min_height, screen_height - 120)

    allowed_max_width = min(screen_max_width, int(max_width)) if max_width else screen_max_width
    allowed_max_height = min(screen_max_height, int(max_height)) if max_height else screen_max_height

    final_width = max(min_width, min(int(width), int(allowed_max_width)))
    final_height = max(min_height, min(int(height), int(allowed_max_height)))

    pos_x = max(0, (screen_width - final_width) // 2)
    pos_y = max(0, (screen_height - final_height) // 3)
    dialog.geometry(f"{final_width}x{final_height}+{pos_x}+{pos_y}")
    return final_width, final_height


def _estimate_message_dialog_size(text, base_width=760, base_height=380):
    lines = text.splitlines() or [""]
    longest_line = max((len(line) for line in lines), default=0)
    line_count = len(lines)

    # Grow width for long lines and height for multiline messages, capped by screen clamp.
    width = base_width + min(320, max(0, longest_line - 70) * 4)
    height = base_height + min(300, max(0, line_count - 10) * 14)
    return width, height


def show_scrollable_text_dialog(title, text):
    dialog = tk.Toplevel()
    dialog.title(title)
    _set_dialog_geometry(dialog, 1000, 620, min_width=760, min_height=420, max_width=920)
    dialog.resizable(True, True)

    frame = tk.Frame(dialog)
    frame.pack(fill="both", expand=True)

    scrollbar = tk.Scrollbar(frame)
    scrollbar.pack(side="right", fill="y")

    text_widget = tk.Text(frame, wrap="none", yscrollcommand=scrollbar.set)
    text_widget.insert("1.0", text)
    text_widget.config(state="disabled")
    text_widget.pack(side="left", fill="both", expand=True)

    scrollbar.config(command=text_widget.yview)

    btn = tk.Button(dialog, text="Close", command=dialog.destroy)
    btn.pack(pady=8)

    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()


def show_focused_info_dialog(title, text, parent=None):
    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title(title)
    est_width, est_height = _estimate_message_dialog_size(text)
    final_width, _ = _set_dialog_geometry(
        dialog,
        est_width,
        est_height,
        min_width=520,
        min_height=260,
        max_width=760,
    )
    dialog.resizable(True, True)

    wrap_length = max(460, min(700, final_width - 40))
    tk.Label(
        dialog,
        text=text,
        padx=16,
        pady=16,
        justify="left",
        anchor="w",
        wraplength=wrap_length,
    ).pack(fill="both", expand=True)

    tk.Button(dialog, text="OK", command=dialog.destroy).pack(pady=(0, 12))

    dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()


def show_focused_yes_no_dialog(title, text, parent=None):
    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title(title)
    est_width, est_height = _estimate_message_dialog_size(text)
    final_width, _ = _set_dialog_geometry(
        dialog,
        est_width,
        est_height,
        min_width=520,
        min_height=260,
        max_width=760,
    )
    dialog.resizable(True, True)

    answer = {"value": False}

    wrap_length = max(460, min(700, final_width - 40))
    tk.Label(
        dialog,
        text=text,
        padx=16,
        pady=16,
        justify="left",
        anchor="w",
        wraplength=wrap_length,
    ).pack(fill="both", expand=True)

    button_row = tk.Frame(dialog)
    button_row.pack(pady=(0, 12))

    def choose_yes():
        answer["value"] = True
        dialog.destroy()

    def choose_no():
        answer["value"] = False
        dialog.destroy()

    tk.Button(button_row, text="Yes", command=choose_yes).pack(side="left", padx=6)
    tk.Button(button_row, text="No", command=choose_no).pack(side="left", padx=6)

    dialog.protocol("WM_DELETE_WINDOW", choose_no)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()
    return answer["value"]






def show_insert_statements_preview():
    # Find the last generated SQL file
    data_dir = ensure_data_dir()
    processed_type = detect_processed_report_type(data_dir)
    tracker_file = data_dir / "last_generated_csv_filename.txt"
    if not tracker_file.exists():
        show_focused_info_dialog("No SQL file", "No last_generated_csv_filename.txt found.")
        return
    csv_stem = tracker_file.read_text(encoding="utf-8").strip().rsplit(".csv", 1)[0]
    sql_file = data_dir / f"{csv_stem}.sql"
    if not sql_file.exists():
        show_focused_info_dialog("No SQL file", f"No SQL file found: {sql_file}")
        return
    sql_text = sql_file.read_text(encoding="utf-8")
    # Extract INSERT statements
    statements = extract_insert_statements(sql_text)
    if not statements:
        show_focused_info_dialog("No INSERTs found", f"No INSERT statements found in {sql_file.name}.")
        return

    # Show in a scrollable dialog
    title = "Preview INSERT Statements"
    if processed_type and processed_type != "UNKNOWN":
        title = f"Preview {processed_type} INSERT Statements"
    show_scrollable_text_dialog(title, "\n\n".join(statements))


def _clean_header_value(value):
    if pd.isna(value):
        return ""

    text = str(value).strip()
    if text.startswith("Unnamed:"):
        return ""

    return text


def _flatten_headers(columns):
    flattened = []
    current_prefix = ""

    for index, column in enumerate(columns, start=1):
        if isinstance(column, tuple):
            parts = [_clean_header_value(part) for part in column]
        else:
            parts = [_clean_header_value(column)]

        if parts and parts[0]:
            current_prefix = parts[0]
        elif current_prefix and parts:

            parts[0] = current_prefix

        combined = []
        for part in parts:
            if part and (not combined or combined[-1] != part):
                combined.append(part)

        flattened.append(" ".join(combined) or f"column_{index}")

    return flattened


def write_db_report(report_text, report_type, data_dir):
    """Append a database report to Step11_db_report.txt with type and timestamp."""
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    report_file = data_dir / "Step11_db_report.txt"
    with open(report_file, "a", encoding="utf-8") as handle:
        handle.write(f"[{report_type}] {timestamp}\n")
        handle.write(report_text.strip() + "\n\n")
    return report_file


def detect_processed_report_type(data_dir):
    """Infer processed report type from last generated CSV filename."""
    tracker_file = data_dir / "last_generated_csv_filename.txt"
    if not tracker_file.exists():
        return "UNKNOWN"

    filename = tracker_file.read_text(encoding="utf-8").strip()
    stem = Path(filename).stem
    match = re.search(
        r"(?<![A-Za-z0-9])(OA|NA|SA|SWOC|EEOC)(?![A-Za-z0-9])",
        stem,
        flags=re.IGNORECASE,
    )
    return match.group(1).upper() if match else "UNKNOWN"




def infer_upload_type(xlsx_file_path=None, csv_file_path=None):
    """Detect staff type from filename or return default."""
    if xlsx_file_path:
        xlsx_staff = _normalize_staff_type(Path(xlsx_file_path).stem)
        if xlsx_staff:
            return xlsx_staff.lower()

    if csv_file_path:
        csv_staff = _normalize_staff_type(Path(csv_file_path).stem)
        if csv_staff:
            return csv_staff.lower()

    return get_config("NEW_OBS_SEM", "DEFAULT_UPLOAD_TYPE", "eeoc")


def get_mysql_connection_settings(section_name):
    db_config = LIVE_CONFIG.get(section_name, {})

    host = db_config.get("DB_HOST")
    user = db_config.get("DB_USER")
    password = db_config.get("DB_PASS")
    database = db_config.get("DB_NAME")

    port_value = db_config.get("DB_PORT") or 3306
    try:
        port = int(port_value)
    except (TypeError, ValueError):
        port = 3306

    ssh_port_value = db_config.get("SSH_PORT", 22)
    try:
        ssh_port = int(ssh_port_value)
    except (TypeError, ValueError):
        ssh_port = 22

    mysql_remote_port_value = db_config.get("MYSQL_REMOTE_PORT", 3306)
    try:
        mysql_remote_port = int(mysql_remote_port_value)
    except (TypeError, ValueError):
        mysql_remote_port = 3306

    ssh_connect_timeout_value = db_config.get("SSH_CONNECT_TIMEOUT", 10)
    try:
        ssh_connect_timeout = int(ssh_connect_timeout_value)
    except (TypeError, ValueError):
        ssh_connect_timeout = 10

    ssh_tunnel_wait_seconds_value = db_config.get("SSH_TUNNEL_WAIT_SECONDS", 20)
    try:
        ssh_tunnel_wait_seconds = float(ssh_tunnel_wait_seconds_value)
    except (TypeError, ValueError):
        ssh_tunnel_wait_seconds = 20.0

    ssh_tunnel_poll_interval_value = db_config.get("SSH_TUNNEL_POLL_INTERVAL", 0.25)
    try:
        ssh_tunnel_poll_interval = float(ssh_tunnel_poll_interval_value)
    except (TypeError, ValueError):
        ssh_tunnel_poll_interval = 0.25

    return {
        "section": section_name,
        "host": host,
        "user": user,
        "password": password,
        "database": database,
        "port": port,
        "ssh_tunnel_enabled": bool(db_config.get("SSH_HOST") and db_config.get("SSH_USER")),
        "ssh_host": db_config.get("SSH_HOST"),
        "ssh_hostname": db_config.get("SSH_HOSTNAME"),
        "ssh_user": db_config.get("SSH_USER"),
        "ssh_port": ssh_port,
        "ssh_key_file": db_config.get("SSH_KEY_FILE") or LIVE_CONFIG.get("SSH_KEY_FILE"),
        "ssh_connect_timeout": ssh_connect_timeout,
        "ssh_tunnel_wait_seconds": ssh_tunnel_wait_seconds,
        "ssh_tunnel_poll_interval": ssh_tunnel_poll_interval,
        "mysql_remote_host": db_config.get("MYSQL_REMOTE_HOST") or "127.0.0.1",
        "mysql_remote_port": mysql_remote_port,
    }






def get_mysql_connection_candidates():
    candidates = []
    settings = get_mysql_connection_settings("DB_SERVER")
    if settings.get("host") and settings.get("user") and settings.get("password"):
        candidates.append(settings)
    return candidates


def connect_to_remote_mysql_db():
    candidates = get_mysql_connection_candidates()
    if not candidates:
        show_focused_info_dialog(
            "Missing MySQL settings",
            "No MySQL connection sections found in common/config.live.ini.\n"
            "Expected section: DB_SERVER."
        )
        return False

    try:
        pymysql = importlib.import_module("pymysql")
    except ImportError:
        show_focused_info_dialog(
            "Missing MySQL client",
            "PyMySQL is not installed. Install it before trying to connect to the remote MySQL database."
        )
        return False

    connect_timeout = get_config("DB_SERVER", "CONNECT_TIMEOUT", 10)
    insert_commit_interval = get_config("DB_SERVER", "INSERT_COMMIT_INTERVAL", 200)
    insert_use_ignore = get_config("DB_SERVER", "INSERT_USE_IGNORE", True)
    insert_detail_limit = get_config("DB_SERVER", "INSERT_DETAIL_LIMIT", 300)
    type_columns = get_config("DB_SERVER", "TYPE_COLUMNS", [])
    staff_types = get_config("DB_SERVER", "STAFF_TYPES", [])
    preferred_tables = get_config("DB_SERVER", "PREFERRED_TABLES", [])

    def _overview_operation(connection, settings):
        db_name = settings["database"]
        overview = collect_mysql_overview(
            connection,
            db_name,
            preferred_tables,
            type_columns,
            staff_types,
        )
        overview["db_name"] = db_name
        return overview

    overview_result = run_with_mysql_connection(
        candidates,
        pymysql,
        _overview_operation,
        connect_timeout=connect_timeout,
        password_provider=get_ssh_password,
    )
    if not overview_result["ok"]:
        show_focused_info_dialog(
            "Database connection failed",
            "Could not connect to the remote MySQL database with any configured profile.\n\n"
            + "\n".join(overview_result["failures"])
        )
        return False

    settings = overview_result["settings"]
    overview = overview_result["result"]
    type_summary = format_type_summary(overview["type_counts"])
    msg = (
        f"Connected using {settings['section']} tunnel {settings['host']}:{settings['port']}\n"
        f"MySQL Server version: {overview['mysql_version']}\n\n"
        f"Database: {overview['db_name']}\nTable: {overview['main_table']}\n\n"
        f"Total rows: {overview['total_rows']}\n"
        + type_summary
    )
    show_focused_info_dialog("Current MySQL Table Stats", msg, parent=getattr(tk, "_default_root", None))
    write_db_report(msg, "BEFORE_INSERT", ensure_data_dir())

    # UI-only behavior stays in Step11 wrapper.
    show_insert_statements_preview()
    apply_choice = show_focused_yes_no_dialog(
        "Apply INSERTs Info",
        "Apply the INSERT statements to the database?\n\n"
        "This will:\n"
        "• Insert new records\n"
        "• Skip records that already exist (duplicate dates)\n"
        "• Report statistics"
    )
    if not apply_choice:
        return True

    data_dir = ensure_data_dir()
    tracker_file = data_dir / "last_generated_csv_filename.txt"
    if not tracker_file.exists():
        return True

    csv_stem = tracker_file.read_text(encoding="utf-8").strip().rsplit(".csv", 1)[0]
    sql_file = data_dir / f"{csv_stem}.sql"
    if not sql_file.exists():
        return True

    def _apply_operation(connection, settings):
        db_name = settings["database"]
        stats = apply_insert_statements_to_database(
            connection,
            sql_file,
            commit_interval=insert_commit_interval,
            use_insert_ignore=insert_use_ignore,
            detail_limit=insert_detail_limit,
        )
        post_overview = collect_mysql_overview(
            connection,
            db_name,
            preferred_tables,
            type_columns,
            staff_types,
        )
        result_file = write_insert_results_to_file(stats, sql_file, data_dir)
        return {
            "stats": stats,
            "post_overview": post_overview,
            "result_file": result_file,
            "main_table": post_overview["main_table"],
        }

    apply_result = run_with_mysql_connection(
        [settings],
        pymysql,
        _apply_operation,
        connect_timeout=connect_timeout,
        password_provider=get_ssh_password,
    )
    if not apply_result["ok"]:
        show_focused_info_dialog(
            "Database connection failed",
            "Could not connect to the remote MySQL database with any configured profile.\n\n"
            + "\n".join(apply_result["failures"])
        )
        return False

    apply_payload = apply_result["result"]
    stats = apply_payload["stats"]
    post_overview = apply_payload["post_overview"]
    processed_type = detect_processed_report_type(data_dir)
    post_type_summary = format_type_summary(post_overview["type_counts"])
    end_delimiter = (
        "\n" + ("=" * 72) + "\n"
        + f"END OF PROCESSING FOR TYPE: {processed_type}\n"
        + ("=" * 72)
    )

    result_msg = (
        f"Database Insert Results for {processed_type}\n\n"
        f"Total statements: {stats['total']}\n"
        f"✓ Inserted successfully: {stats['success']}\n"
        f"⊘ Already existed (skipped): {stats['already_exists']}\n"
        f"✗ Failed: {stats['failed']}\n\n"
        f"After Insert Snapshot\n"
        f"Table: {apply_payload['main_table']}\n\n"
        f"Total rows: {post_overview['total_rows']}\n"
        f"{post_type_summary}\n\n"
        f"Detailed results saved to:\n{apply_payload['result_file'].name}\n"
        f"{end_delimiter}"
    )

    show_focused_info_dialog("Insert Results", result_msg)
    write_db_report(result_msg, "AFTER_INSERT", ensure_data_dir())
    return True


def preflight_mysql_connection_check():
    candidates = get_mysql_connection_candidates()
    if not candidates:
        show_focused_info_dialog(
            "Missing MySQL settings",
            "No MySQL connection sections found in common/config.live.ini.\n"
            "Expected section: DB_SERVER."
        )
        return False

    try:
        pymysql = importlib.import_module("pymysql")
    except ImportError:
        show_focused_info_dialog(
            "Missing MySQL client",
            "PyMySQL is not installed. Install it before trying database import."
        )
        return False

    query_timeout = get_config("DB_SERVER", "QUERY_TIMEOUT", 5)
    result = run_mysql_preflight(
        candidates,
        pymysql,
        query_timeout=query_timeout,
        password_provider=get_ssh_password,
    )
    if result["ok"]:
        return True

    show_focused_info_dialog(
        "Database preflight failed",
        "Database authentication failed before import.\n\n"
        "Fix credentials/grants, then rerun this step.\n\n"
        + "\n".join(result["failures"])
    )
    return False


def upload_csv_to_staff_site(csv_path, upload_type, verbose_enabled):
    """Upload CSV file to staff schedule service."""
    session = requests.Session()
    session.mount("https://", LegacyTLSAdapter())

    form_data = {"type": upload_type, "submit": "Submit"}
    if verbose_enabled:
        form_data["verbose"] = "on"

    upload_url = get_config("NEW_OBS_SEM", "STAFF_UPLOAD_URL")
    with open(csv_path, "rb") as file_handle:
        response = session.post(
            str(upload_url),
            data=form_data,
            files={"file": (Path(csv_path).name, file_handle, "text/csv")},
            timeout=60,
        )

    response.raise_for_status()
    return response


def open_file_in_integrated_browser(output_file):
    # Keep focus in Finder on macOS instead of foregrounding VS Code.
    if sys.platform == "darwin":
        try:
            subprocess.Popen(["open", "-R", str(output_file)])
            return output_file
        except Exception:
            subprocess.Popen(["open", str(output_file)])
            return output_file

    return output_file


def open_response_in_integrated_browser(response_text, output_file):
    output_file.write_text(response_text, encoding="utf-8")
    return open_file_in_integrated_browser(output_file)






def pick_file(root):
    picker_parent = tk.Toplevel(root)
    picker_parent.overrideredirect(True)
    picker_parent.geometry("1x1+0+0")
    picker_parent.attributes("-alpha", 0.0)
    picker_parent.attributes("-topmost", True)
    picker_parent.lift()
    picker_parent.focus_force()
    picker_parent.update_idletasks()

    try:
        return filedialog.askopenfilename(
            parent=picker_parent,
            title="Select Excel file",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
    finally:
        try:
            picker_parent.destroy()
        except Exception:
            pass

def get_sheet_names(file_path):
    start_time = time.perf_counter()
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        elapsed_ms = (time.perf_counter() - start_time) * 1000.0
        _debug_log(
            f"Sheet name discovery took {elapsed_ms:.1f} ms for {Path(file_path).name} "
            f"({len(sheet_names)} sheets)"
        )
        return sheet_names
    finally:
        workbook.close()

def choose_sheet(sheet_names, file_path, default_sheet=None, parent=None):
    dialog = tk.Toplevel(parent)
    dialog.title("Select Sheet to Convert")
    _set_dialog_geometry(dialog, 820, 360, min_width=660, min_height=300, max_width=760)
    dialog.resizable(True, True)

    default_name = default_sheet or sheet_names[0]
    selected = tk.StringVar(master=dialog, value=default_name)
    cancelled = [False]

    tk.Label(
        dialog,
        text=(
            f"Excel (xlsx) file:\n{file_path}\n\n"
            f"Accept default or select Sheet name from the dropdown list, below\n\n"
            f"Then OK to continue..."
        ),
        padx=16,
        pady=16,
        justify="left",
        wraplength=660
    ).pack()

    dropdown = tk.OptionMenu(dialog, selected, *sheet_names)
    dropdown.pack(pady=(0, 8))

    def confirm():
        dialog.destroy()

    def cancel():
        cancelled[0] = True
        dialog.destroy()

    btn_frame = tk.Frame(dialog)
    btn_frame.pack(pady=10)
    tk.Button(btn_frame, text="OK", command=confirm).pack(side="left", padx=6)
    tk.Button(btn_frame, text="Cancel", command=cancel).pack(side="left", padx=6)

    dialog.protocol("WM_DELETE_WINDOW", cancel)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()
    return None if cancelled[0] else selected.get()

def confirm_staff_type_dialog(staff_type, file_path, parent=None):
    """Show detected staff type and let user confirm before proceeding."""
    dialog = tk.Toplevel(parent)
    dialog.title("Confirm Staff Type")
    _set_dialog_geometry(dialog, 560, 280, min_width=500, min_height=240)
    dialog.resizable(True, True)

    staff_type_display = staff_type.upper() if staff_type else "UNKNOWN"
    cancelled = [False]

    tk.Label(
        dialog,
        text=(
            f"File: {Path(file_path).name}\n\n"
            f"Detected staff type: {staff_type_display}\n\n"
            "Is this correct?"
        ),
        padx=16,
        pady=20,
        justify="left",
        wraplength=460
    ).pack()

    def confirm():
        dialog.destroy()

    def cancel():
        cancelled[0] = True
        dialog.destroy()

    btn_frame = tk.Frame(dialog)
    btn_frame.pack(pady=10)
    tk.Button(btn_frame, text="Yes, Continue", command=confirm).pack(side="left", padx=6)
    tk.Button(btn_frame, text="No, Cancel", command=cancel).pack(side="left", padx=6)

    dialog.protocol("WM_DELETE_WINDOW", cancel)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()
    return not cancelled[0]


def apply_staff_type_processing(df, staff_type):
    """Apply type-specific data transformations and validations."""
    if not staff_type:
        return df

    staff_type_lower = staff_type.lower()

    # Type-specific processing
    if staff_type_lower == "na":
        # NA-specific: ensure Date column exists (schedule format)
        if "Date" not in df.columns:
            show_focused_info_dialog(
                "Invalid NA file",
                "NA schedule files must contain a 'Date' column."
            )
    elif staff_type_lower == "oa":
        # OA-specific: validation/processing (to be defined)
        pass
    elif staff_type_lower == "sa":
        # SA-specific: validation/processing (to be defined)
        pass

    return df

def open_in_calc(file_path, sheet_name):
    launcher = shutil.which("libreoffice") or shutil.which("soffice")
    if not launcher:
        show_focused_info_dialog(
            "LibreOffice not found",
            "Could not find LibreOffice.\nInstall LibreOffice or add it to PATH."
        )
        return None

    return subprocess.Popen([
        launcher,
        "--norestore",
        "--calc",
        file_path
    ])


def close_calc_process(process):
    if process is None:
        return

    # On macOS, ask the app to quit first to avoid crash-recovery prompts.
    if sys.platform == "darwin":
        try:
            subprocess.run(
                ["osascript", "-e", 'tell application "LibreOffice" to quit'],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=3,
            )
            if process.poll() is None:
                process.wait(timeout=8)
            return
        except Exception:
            pass

    if process.poll() is not None:
        return

    process.terminate()
    try:
        process.wait(timeout=8)
    except subprocess.TimeoutExpired:
        # Do not force-kill; hard kills can trigger LibreOffice recovery on next launch.
        return


def focus_converted_file(file_path):
    if sys.platform == "darwin":
        target = Path(file_path)
        if target.is_dir():
            subprocess.Popen(["open", str(target)])
        else:
            subprocess.Popen(["open", "-R", str(target)])


def show_final_completion_dialog(output=None, upload_message=None, sql_file=None, parent=None):
    dialog = tk.Toplevel(parent)
    dialog.title("Conversions Complete")
    dialog.resizable(True, True)

    action = [""]

    def continue_to_database():
        action[0] = "ok"
        dialog.destroy()

    def cancel():
        action[0] = "cancel"
        dialog.destroy()

    button_frame = tk.Frame(dialog)
    button_frame.pack(side="bottom", fill="x", pady=(0, 12))
    tk.Button(button_frame, text="OK", command=continue_to_database).pack(side="left", padx=6)
    tk.Button(button_frame, text="Cancel", command=cancel).pack(side="left", padx=6)

    export_section = ""
    if output and upload_message:
        export_section = f"{upload_message}\n\n"
        export_section += f"Converted CSV file:\n{output}\n\n"
        if sql_file:
            export_section += f"Saved SQL file:\n{sql_file}\n\nand is ready for database update.\n\n"

    est_width, est_height = _estimate_message_dialog_size(export_section or "Conversions complete.")
    final_width, _ = _set_dialog_geometry(
        dialog,
        est_width,
        est_height,
        min_width=620,
        min_height=320,
        max_width=760,
    )

    if export_section:
        wrap_length = max(520, min(700, final_width - 40))
        tk.Label(
            dialog,
            text=export_section.rstrip(),
            padx=16,
            pady=20,
            justify="left",
            wraplength=wrap_length
        ).pack()

    # When completion is shown, immediately reveal/select the generated SQL file.
    if sql_file:
        focus_converted_file(sql_file)

    dialog.protocol("WM_DELETE_WINDOW", cancel)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()
    return action[0]


def export_sheet(file_path, sheet_name, data_dir, upload_enabled, staff_type=None):
    df = pd.read_excel(
        file_path,
        sheet_name=sheet_name,
        header=[0, 1],
        engine="openpyxl"
    )
    df.columns = _flatten_headers(df.columns)

    # Apply staff-type-specific processing
    df = apply_staff_type_processing(df, staff_type)

    if "Date" in df.columns:
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce").dt.strftime("%m/%d/%y")

    output = data_dir / build_output_filename(file_path)
    df.to_csv(output, index=False)
    store_last_csv_filename(output, data_dir)
    sql_file = None

    if upload_enabled:
        upload_url = get_config("NEW_OBS_SEM", "STAFF_UPLOAD_URL")
        _debug_log(
            "Upload URL check: "
            f"value={repr(upload_url)}, "
            f"new_obs_sem_type={type(LIVE_CONFIG.get('NEW_OBS_SEM')).__name__}, "
            f"config_sections={list(LIVE_CONFIG.keys())}"
        )
        if not upload_url:
            show_focused_info_dialog(
                "Invalid upload URL",
                "STAFF_UPLOAD_URL is missing in common/config.live.ini"
            )
            return False

        upload_type = infer_upload_type(xlsx_file_path=file_path, csv_file_path=output)
        try:
            verbose_response = upload_csv_to_staff_site(output, upload_type, verbose_enabled=True)
            verbose_file = data_dir / "last_upload_verbose.html"
            open_response_in_integrated_browser(verbose_response.text, verbose_file)

            final_response = upload_csv_to_staff_site(output, upload_type, verbose_enabled=False)
            final_results_file = data_dir / "last_upload_sql_results.html"
            open_response_in_integrated_browser(final_response.text, final_results_file)
            sql_file = save_insert_statements(final_response.text, output, data_dir)
            upload_message = (
                f"===== Processing staff type: {upload_type.upper()} =====\n"
                #f"Verbose HTTP {verbose_response.status_code}\nFinal HTTP {final_response.status_code}\n\n"
            )
        except Exception as error:
            upload_message = f"Upload failed: {error}"
    else:
        upload_message = "Upload skipped (Automatic upload is disabled)."

    return True, output, upload_message, sql_file


# ============================================================================
# SWOC Rotation Management Functions
# ============================================================================

def load_swoc_members():
    """Load SWOC team members from config."""
    members = get_config("SWOC_ROTATION", "MEMBERS", [])
    _debug_log(
        "SWOC members load: "
        f"type={type(members).__name__}, "
        f"count={(len(members) if isinstance(members, list) else 0)}, "
        f"config_sections={list(LIVE_CONFIG.keys())}"
    )
    if not isinstance(members, list):
        return []

    normalized = []
    for member in members:
        if not isinstance(member, dict):
            continue
        # Backward compatibility: migrate legacy "initials" to "alias" in-memory.
        if "alias" not in member and "initials" in member:
            member = dict(member)
            member["alias"] = member.get("initials")
            member.pop("initials", None)
        normalized.append(member)

    return normalize_rotation_order(normalized)


def _coerce_positive_int(value):
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return None
    return parsed if parsed > 0 else None


def normalize_rotation_order(members, preserve_list_order=False):
    """Rewrite rotation_order to contiguous 1..N, optionally preserving list order."""
    if preserve_list_order:
        normalized = []
        for position, member in enumerate(members, start=1):
            member_copy = dict(member)
            member_copy["rotation_order"] = position
            normalized.append(member_copy)
        return normalized

    sortable = []
    for original_index, member in enumerate(members):
        order_value = _coerce_positive_int(member.get("rotation_order"))
        fallback_order = original_index + 1
        sortable.append((order_value if order_value is not None else fallback_order, original_index, dict(member)))

    sortable.sort(key=lambda row: (row[0], row[1]))

    normalized = []
    for position, (_, _, member) in enumerate(sortable, start=1):
        member["rotation_order"] = position
        normalized.append(member)
    return normalized


def save_swoc_members(members):
    """Save SWOC team members back to config file by patching only the MEMBERS list."""
    config_file = PROJECT_ROOT / "common" / "config.live.ini"
    if not config_file.exists():
        return False

    content = config_file.read_text(encoding="utf-8")

    # Render the new MEMBERS list as indented Python literals
    indent = "    "
    item_lines = []
    for m in members:
        item_lines.append(f"{indent}  {{")
        for k, v in m.items():
            item_lines.append(f'{indent}    "{k}": {repr(v)},')
        # Remove trailing comma from last key
        item_lines[-1] = item_lines[-1].rstrip(",")
        item_lines.append(f"{indent}  }},")
    if item_lines:
        item_lines[-1] = item_lines[-1].rstrip(",")
    members_block = "\n".join(item_lines)

    new_section = f'{indent}"MEMBERS": [\n{members_block}\n{indent}]'

    # Replace existing MEMBERS block using a regex that finds "MEMBERS": [ ... ]
    import re as _re
    pattern = _re.compile(
        r'"MEMBERS"\s*:\s*\[.*?\]',
        _re.DOTALL,
    )
    if pattern.search(content):
        new_content = pattern.sub(new_section.strip(), content, count=1)
    else:
        # Append before closing brace of SWOC_ROTATION if MEMBERS key is missing
        new_content = content.rstrip().rstrip("}").rstrip() + f',\n{new_section}\n  }}\n}}\n'

    config_file.write_text(new_content, encoding="utf-8")

    global LIVE_CONFIG
    LIVE_CONFIG = load_live_config()
    return True


def write_swoc_member_list_log(members, changed=True):
    """Write SWOC member list to log file."""
    data_dir = ensure_data_dir()
    log_file = data_dir / "swoc_member_list.log"

    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    header = f"[{timestamp}] SWOC Member List\n"

    if changed:
        log_lines = [
            header,
            "  The new member rotation order has been changed.",
            "  New rotation order:",
        ]
        for i, member in enumerate(members, start=1):
            line = (
                f"  #{i:02d} {member.get('name')} ({member.get('alias', '')})"
            )
            log_lines.append(line)
        log_lines.append("\n")
    else:
        log_lines = [
            header,
            "  No change to member rotation order was made.\n",
        ]

    with open(log_file, "a", encoding="utf-8") as f:
        f.write("\n".join(log_lines))

    return log_file


# ---------------------------------------------------------------------------
# SWOC schedule generation helpers
# ---------------------------------------------------------------------------

def _get_upcoming_semester_info():
    """Return info dict for the next upcoming semester relative to today.

    Returns a dict: {"label": str, "start": date, "end": date}.
    """
    today = datetime.date.today()
    sem_a_cfg = get_config("SWOC_ROTATION", "SEMESTER_A", {})
    sem_b_cfg = get_config("SWOC_ROTATION", "SEMESTER_B", {})

    sem_a_start = sem_a_end = sem_a_label = None
    sem_b_start = sem_b_end = sem_b_label = None

    try:
        sem_a_start = datetime.date.fromisoformat(sem_a_cfg.get("start_date", ""))
        sem_a_end = datetime.date.fromisoformat(sem_a_cfg.get("end_date", ""))
        sem_a_label = f"{sem_a_start.year}A"
    except (ValueError, AttributeError, TypeError):
        pass

    try:
        sem_b_start = datetime.date.fromisoformat(sem_b_cfg.get("start_date", ""))
        sem_b_end = datetime.date.fromisoformat(sem_b_cfg.get("end_date", ""))
        sem_b_label = f"{sem_b_start.year}B"
    except (ValueError, AttributeError, TypeError):
        pass

    candidates = []
    if sem_a_start and sem_a_start > today and sem_a_end:
        candidates.append({"label": sem_a_label, "start": sem_a_start, "end": sem_a_end})
    if sem_b_start and sem_b_start > today and sem_b_end:
        candidates.append({"label": sem_b_label, "start": sem_b_start, "end": sem_b_end})

    if candidates:
        candidates.sort(key=lambda x: x["start"])
        return candidates[0]

    # Both semesters in past or current — return B as fallback
    if sem_b_start and sem_b_end:
        return {"label": sem_b_label or "Unknown", "start": sem_b_start, "end": sem_b_end}
    return {"label": "Unknown", "start": today, "end": today}


def _reveal_in_finder_and_refocus(path, refocus_widget):
    """Reveal file in Finder (macOS) and schedule focus return to refocus_widget."""
    if sys.platform == "darwin":
        try:
            subprocess.Popen(["open", "-R", str(path)])
        except Exception:
            pass
        if refocus_widget:
            refocus_widget.after(400, lambda: (
                refocus_widget.lift(),
                refocus_widget.focus_force(),
            ))


def _show_log_saved_and_prompt(log_file, parent):
    """Show 'log saved' notification with Close / Continue buttons.

    Returns 'close' or 'continue'.
    """
    owner = parent or getattr(tk, "_default_root", None)
    notify = tk.Toplevel(owner)
    notify.title("SWOC Member Log Saved")
    _set_dialog_geometry(notify, 700, 240, min_width=540, min_height=200)
    notify.resizable(True, True)

    result = {"action": "close"}

    msg = (
        "SWOC member list log was updated.\n\n"
        f"Log file:  {log_file.name}\n"
        f"Path:  {log_file}\n\n"
        "Review the log file in Finder if desired, then choose an action below."
    )
    final_width, _ = _set_dialog_geometry(notify, 700, 260, min_width=540, min_height=220)
    wrap_length = max(480, final_width - 40)
    tk.Label(
        notify,
        text=msg,
        padx=16,
        pady=16,
        justify="left",
        anchor="w",
        wraplength=wrap_length,
    ).pack(fill="both", expand=True)

    btn_frame = tk.Frame(notify)
    btn_frame.pack(pady=(0, 12))

    def on_close():
        result["action"] = "close"
        notify.destroy()

    def on_continue():
        result["action"] = "continue"
        notify.destroy()

    tk.Button(btn_frame, text="Close", command=on_close, width=10).pack(side="left", padx=8)
    tk.Button(btn_frame, text="Continue →", command=on_continue, width=18).pack(side="left", padx=8)

    notify.protocol("WM_DELETE_WINDOW", on_close)
    notify.lift()
    notify.focus_force()
    notify.attributes("-topmost", True)
    notify.grab_set()
    notify.wait_window()

    return result["action"]


def show_swoc_date_range_dialog(parent, default_start, default_end, semester_label):
    """Prompt for start/end dates for SWOC schedule generation.

    Returns (start_date, end_date) or None if cancelled.
    """
    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title(f"SWOC Schedule Date Range — {semester_label}")
    _set_dialog_geometry(dialog, 460, 230, min_width=400, min_height=210)
    dialog.resizable(False, False)

    result_dates = {
        "start": default_start,
        "end": default_end,
        "updated": False,
    }

    frame = tk.Frame(dialog, padx=16, pady=14)
    frame.pack(fill="both", expand=True)

    tk.Label(frame, text=f"Semester: {semester_label}", font=("Arial", 11, "bold")).grid(
        row=0, column=0, columnspan=2, sticky="w", pady=(0, 10)
    )

    tk.Label(frame, text="Start date (YYYY-MM-DD):").grid(row=1, column=0, sticky="w", pady=4)
    start_entry = tk.Entry(frame, width=16)
    start_entry.insert(0, default_start.isoformat())
    start_entry.grid(row=1, column=1, sticky="w", pady=4, padx=(8, 0))

    tk.Label(frame, text="End date (YYYY-MM-DD):").grid(row=2, column=0, sticky="w", pady=4)
    end_entry = tk.Entry(frame, width=16)
    end_entry.insert(0, default_end.isoformat())
    end_entry.grid(row=2, column=1, sticky="w", pady=4, padx=(8, 0))

    btn_frame = tk.Frame(frame)
    btn_frame.grid(row=3, column=0, columnspan=2, pady=(14, 0))

    def on_ok():
        try:
            start = datetime.date.fromisoformat(start_entry.get().strip())
            end = datetime.date.fromisoformat(end_entry.get().strip())
        except ValueError:
            show_focused_info_dialog("Invalid Date", "Please enter dates in YYYY-MM-DD format.", dialog)
            return
        if end < start:
            show_focused_info_dialog("Invalid Range", "End date must be on or after start date.", dialog)
            return
        result_dates["start"] = start
        result_dates["end"] = end
        result_dates["updated"] = True
        dialog.destroy()

    def on_cancel():
        dialog.destroy()

    tk.Button(btn_frame, text="OK", command=on_ok, width=8).pack(side="left", padx=6)
    tk.Button(btn_frame, text="Cancel", command=on_cancel, width=8).pack(side="left", padx=6)

    dialog.protocol("WM_DELETE_WINDOW", on_cancel)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()

    if result_dates["updated"]:
        return result_dates["start"], result_dates["end"]
    return None


def _write_swoc_generation_log(
    semester_label,
    start_date,
    end_date,
    last_alias,
    last_date,
    start_index,
    active_members,
    statements_count,
):
    """Write SWOC generation details and query info to a log file."""
    data_dir = ensure_data_dir()
    log_file = data_dir / "swoc_schedule_generation.log"
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    lines = [
        f"[{timestamp}] SWOC Schedule Generation",
        f"Semester: {semester_label}",
        f"Date range: {start_date} to {end_date}",
        "",
        "Query used:",
        "SELECT Alias, Date FROM nightStaff WHERE Type = 'swoc' ORDER BY Date DESC LIMIT 1;",
        f"Last swoc alias/date: {last_alias or 'N/A'} / {last_date or 'N/A'}",
        f"Starting index after advance: {start_index}",
        (
            "Starting alias after advance: "
            f"{active_members[start_index].get('alias') if active_members else 'N/A'}"
        ),
        f"Total generated INSERT statements: {statements_count}",
        "",
    ]

    with open(log_file, "a", encoding="utf-8") as f:
        f.write("\n".join(lines))

    return log_file


def _generate_swoc_insert_statements(members, start_date, end_date, start_index):
    """Generate one INSERT per day cycling through active members.

    Returns list of SQL strings.
    """
    if not members:
        return []

    statements = []
    current = start_date
    idx = start_index % len(members)
    one_day = datetime.timedelta(days=1)

    while current <= end_date:
        alias = members[idx].get("alias", "")
        date_str = current.isoformat()
        stmt = (
            f"insert into nightStaff set "
            f"Date='{date_str}', TelNr='0', Alias='{alias}', Type='swoc';"
        )
        statements.append(stmt)
        idx = (idx + 1) % len(members)
        current += one_day

    return statements


def _show_swoc_sql_preview_dialog(parent, sql_file, semester_label):
    """Show SQL file path with Exit / Publish buttons. Returns 'exit' or 'publish'."""
    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title(f"SWOC Schedule SQL — {semester_label}")
    _set_dialog_geometry(dialog, 720, 280, min_width=560, min_height=240)
    dialog.resizable(True, True)

    result = {"action": "exit"}

    msg = (
        f"SWOC INSERT statements have been written.\n\n"
        f"SQL file:  {sql_file.name}\n"
        f"Path:  {sql_file}\n\n"
        "Review the file in Finder before publishing.\n"
        "Press Exit to stop here, or Publish to insert into keckOperations."
    )
    final_width, _ = _set_dialog_geometry(dialog, 720, 280, min_width=560, min_height=240)
    wrap_length = max(500, final_width - 40)
    tk.Label(
        dialog,
        text=msg,
        padx=16,
        pady=16,
        justify="left",
        anchor="w",
        wraplength=wrap_length,
    ).pack(fill="both", expand=True)

    btn_frame = tk.Frame(dialog)
    btn_frame.pack(pady=(0, 12))

    def on_exit():
        result["action"] = "exit"
        dialog.destroy()

    def on_publish():
        result["action"] = "publish"
        dialog.destroy()

    tk.Button(btn_frame, text="Exit", command=on_exit, width=10).pack(side="left", padx=8)
    tk.Button(
        btn_frame,
        text=f"Publish SWOC schedule for semester {semester_label}",
        command=on_publish,
    ).pack(side="left", padx=8)

    dialog.protocol("WM_DELETE_WINDOW", on_exit)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()

    return result["action"]


def _show_swoc_results_dialog(parent, stats, result_file, semester_label):
    """Show INSERT results summary; OK button opens the results file."""
    total = stats.get("total", 0)
    success = stats.get("success", 0)
    already_exists = stats.get("already_exists", 0)
    failed = stats.get("failed", 0)

    # Collect per-row details for failures and duplicates
    failed_lines = []
    skipped_lines = []
    for stmt, status, error in stats.get("statements", []):
        date_m = re.search(r"Date='([^']+)'", stmt)
        alias_m = re.search(r"Alias='([^']+)'", stmt)
        label = (
            f"{date_m.group(1) if date_m else '?'} / "
            f"{alias_m.group(1) if alias_m else '?'}"
        )
        if status == "FAILED":
            err_snippet = (error[:60] if error else "")
            failed_lines.append(f"  \u2717 {label}  [{err_snippet}]")
        elif status == "DUPLICATE":
            skipped_lines.append(f"  \u2298 {label}")

    parts = [
        f"SWOC Schedule Insert Results \u2014 {semester_label}\n",
        f"Total attempted:           {total}",
        f"\u2713 Inserted successfully:   {success}",
        f"\u2298 Already existed (skip):  {already_exists}",
        f"\u2717 Failed:                  {failed}",
    ]
    if skipped_lines:
        parts.append(f"\nSkipped ({min(len(skipped_lines), 20)} of {len(skipped_lines)} shown):")
        parts.extend(skipped_lines[:20])
        if len(skipped_lines) > 20:
            parts.append(f"  \u2026 and {len(skipped_lines) - 20} more (see results file)")
    if failed_lines:
        parts.append(f"\nFailed ({min(len(failed_lines), 20)} of {len(failed_lines)} shown):")
        parts.extend(failed_lines[:20])
        if len(failed_lines) > 20:
            parts.append(f"  \u2026 and {len(failed_lines) - 20} more (see results file)")
    parts.append(f"\nResults file:  {result_file.name}")
    parts.append(f"Path:  {result_file}")

    msg = "\n".join(parts)

    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title(f"SWOC Publish Results \u2014 {semester_label}")
    est_w, est_h = _estimate_message_dialog_size(msg)
    _set_dialog_geometry(dialog, max(est_w, 700), max(est_h, 360), min_width=580, min_height=320)
    dialog.resizable(True, True)

    outer = tk.Frame(dialog)
    outer.pack(fill="both", expand=True)

    sb = tk.Scrollbar(outer)
    sb.pack(side="right", fill="y")
    txt = tk.Text(outer, wrap="word", yscrollcommand=sb.set, padx=12, pady=12)
    txt.insert("1.0", msg)
    txt.config(state="disabled")
    txt.pack(side="left", fill="both", expand=True)
    sb.config(command=txt.yview)

    def on_ok():
        dialog.destroy()
        if sys.platform == "darwin":
            try:
                subprocess.Popen(["open", str(result_file)])
            except Exception:
                pass

    tk.Button(dialog, text="OK \u2014 View Results", command=on_ok).pack(pady=(0, 12))

    dialog.protocol("WM_DELETE_WINDOW", on_ok)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()


def run_swoc_schedule_generation(parent=None):
    """Orchestrate the full SWOC schedule generation and publish flow."""
    # 1. Upcoming semester defaults
    sem_info = _get_upcoming_semester_info()
    semester_label = sem_info["label"]
    default_start = sem_info["start"]
    default_end = sem_info["end"]

    # 2. Date range dialog
    dates = show_swoc_date_range_dialog(parent, default_start, default_end, semester_label)
    if not dates:
        return
    start_date = dates[0]
    end_date = dates[1]

    # 3. Active members sorted by rotation_order
    all_members = load_swoc_members()
    active_members = [m for m in all_members if m.get("active", True)]
    if not active_members:
        show_focused_info_dialog(
            "No Active Members",
            "No active SWOC rotation members found. Add or activate members before generating a schedule.",
            parent,
        )
        return

    # 4. Query last SWOC entry to determine starting index
    start_index = 0
    last_alias = None
    last_date = None
    candidates = get_mysql_connection_candidates()
    if candidates:
        try:
            pymysql = importlib.import_module("pymysql")
            connect_timeout = get_config("DB_SERVER", "CONNECT_TIMEOUT", 10)

            def _query_last_swoc(connection, _settings):
                with connection.cursor() as cursor:
                    cursor.execute(
                        "SELECT Alias, Date FROM nightStaff "
                        "WHERE Type = 'swoc' ORDER BY Date DESC LIMIT 1"
                    )
                    return cursor.fetchone()  # (alias, date) or None

            query_result = run_with_mysql_connection(
                candidates,
                pymysql,
                _query_last_swoc,
                connect_timeout=connect_timeout,
                password_provider=get_ssh_password,
            )
            if query_result["ok"] and query_result["result"]:
                last_alias, last_date = query_result["result"]
                alias_to_idx = {m.get("alias"): i for i, m in enumerate(active_members)}
                if last_alias in alias_to_idx:
                    last_idx = alias_to_idx[last_alias]
                    start_index = (last_idx + 1) % len(active_members)
                    _debug_log(
                        f"Last SWOC entry: alias={last_alias}, date={last_date}. "
                        f"Starting rotation from index {start_index} "
                        f"({active_members[start_index].get('alias')})"
                    )
        except Exception as exc:
            _debug_log(f"Could not query last SWOC entry (non-fatal): {exc}")

    # 5. Generate INSERT statements
    statements = _generate_swoc_insert_statements(active_members, start_date, end_date, start_index)
    if not statements:
        show_focused_info_dialog(
            "No Statements Generated",
            f"No INSERT statements were generated for {start_date} to {end_date}.",
            parent,
        )
        return

    # 6. Write SQL file
    data_dir = ensure_data_dir()
    sql_file = data_dir / f"{semester_label}_SWOC_Schedule.sql"
    sql_file.write_text("\n".join(statements) + "\n", encoding="utf-8")
    _debug_log(f"Wrote {len(statements)} SWOC INSERT statements to {sql_file}")

    generation_log = _write_swoc_generation_log(
        semester_label=semester_label,
        start_date=start_date,
        end_date=end_date,
        last_alias=last_alias,
        last_date=last_date,
        start_index=start_index,
        active_members=active_members,
        statements_count=len(statements),
    )

    # Reveal generation log in Finder first (includes query details), then return focus.
    _reveal_in_finder_and_refocus(generation_log, parent)

    # 7. Reveal SQL file in Finder; return focus to parent
    _reveal_in_finder_and_refocus(sql_file, parent)

    # 8. SQL preview dialog — Exit or Publish
    action = _show_swoc_sql_preview_dialog(parent, sql_file, semester_label)
    if action != "publish":
        return

    # 9. Publish to database
    candidates = get_mysql_connection_candidates()
    if not candidates:
        show_focused_info_dialog(
            "No DB Config",
            "No MySQL connection settings found. Configure DB_SERVER in config.live.ini.",
            parent,
        )
        return

    try:
        pymysql = importlib.import_module("pymysql")
    except ImportError:
        show_focused_info_dialog(
            "Missing pymysql",
            "PyMySQL is not installed. Install it to publish to the database.",
            parent,
        )
        return

    insert_commit_interval = get_config("DB_SERVER", "INSERT_COMMIT_INTERVAL", 200)
    insert_use_ignore = get_config("DB_SERVER", "INSERT_USE_IGNORE", True)
    insert_detail_limit = get_config("DB_SERVER", "INSERT_DETAIL_LIMIT", 300)

    def _publish_operation(connection, _settings):
        stats = apply_insert_statements_to_database(
            connection,
            sql_file,
            commit_interval=insert_commit_interval,
            use_insert_ignore=insert_use_ignore,
            detail_limit=insert_detail_limit,
        )
        result_file = write_insert_results_to_file(stats, sql_file, data_dir)
        return {"stats": stats, "result_file": result_file}

    publish_result = run_with_mysql_connection(
        candidates,
        pymysql,
        _publish_operation,
        connect_timeout=get_config("DB_SERVER", "CONNECT_TIMEOUT", 10),
        password_provider=get_ssh_password,
    )
    if not publish_result["ok"]:
        show_focused_info_dialog(
            "Database connection failed",
            "Could not connect to keckOperations.\n\n" + "\n".join(publish_result["failures"]),
            parent,
        )
        return

    payload = publish_result["result"]
    stats = payload["stats"]
    result_file = payload["result_file"]

    # 10. Reveal results in Finder, show summary
    _reveal_in_finder_and_refocus(result_file, parent)
    _show_swoc_results_dialog(parent, stats, result_file, semester_label)


# ---------------------------------------------------------------------------
# End SWOC schedule generation helpers
# ---------------------------------------------------------------------------


def add_swoc_audit_log_entry(change_type, member_name, old_values, new_values, reason=""):
    """Write an entry to the SWOC rotation audit log."""
    data_dir = ensure_data_dir()
    audit_file = data_dir / "swoc_rotation_audit.log"
    
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    entry = (
        f"[{timestamp}] {change_type}\n"
        f"  Member: {member_name}\n"
    )
    
    if old_values:
        entry += f"  Old values: {old_values}\n"
    if new_values:
        entry += f"  New values: {new_values}\n"
    if reason:
        entry += f"  Reason: {reason}\n"
    
    entry += "\n"
    
    with open(audit_file, "a", encoding="utf-8") as f:
        f.write(entry)


def show_member_edit_dialog(member_name, member_data, parent=None):
    """
    Show a form to edit a SWOC member's lifecycle dates.
    Returns updated member dict or None if cancelled.
    """
    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title(f"Edit Member: {member_name}")
    _set_dialog_geometry(dialog, 520, 450, min_width=440, min_height=380)
    dialog.resizable(False, False)
    
    result = {"updated": False, "data": None}
    
    # Create a frame with padding
    main_frame = tk.Frame(dialog, padx=12, pady=12)
    main_frame.pack(fill="both", expand=True)
    
    # Helper to create labeled entry fields
    fields = {}
    
    def create_date_field(label_text, initial_value, row):
        tk.Label(main_frame, text=label_text + ":").grid(row=row, column=0, sticky="w", pady=4)
        entry = tk.Entry(main_frame, width=18)
        entry.insert(0, initial_value or "")
        entry.grid(row=row, column=1, sticky="ew", pady=4, padx=(8, 0))
        return entry
    
    def create_bool_field(label_text, initial_value, row, command=None):
        tk.Label(main_frame, text=label_text + ":").grid(row=row, column=0, sticky="w", pady=4)
        var = tk.BooleanVar(value=initial_value)
        if command is None:
            check = tk.Checkbutton(main_frame, variable=var)
        else:
            check = tk.Checkbutton(main_frame, variable=var, command=command)
        check.grid(row=row, column=1, sticky="w", pady=4, padx=(8, 0))
        return var

    def create_text_field(label_text, initial_value, row):
        tk.Label(main_frame, text=label_text + ":").grid(row=row, column=0, sticky="w", pady=4)
        entry = tk.Entry(main_frame, width=18)
        entry.insert(0, "" if initial_value is None else str(initial_value))
        entry.grid(row=row, column=1, sticky="ew", pady=4, padx=(8, 0))
        return entry
    
    # Create fields
    fields["service_start_date"] = create_date_field(
        "Service Start Date (YYYY-MM-DD)", 
        member_data.get("service_start_date"), 0
    )
    fields["service_end_date"] = create_date_field(
        "Service End Date (YYYY-MM-DD)", 
        member_data.get("service_end_date"), 1
    )
    fields["rotation_order"] = create_text_field(
        "Rotation Order #",
        member_data.get("rotation_order"), 2
    )
    fields["rotation_start_date"] = create_date_field(
        "Rotation Start Date (YYYY-MM-DD)", 
        member_data.get("rotation_start_date"), 4
    )
    fields["rotation_end_date"] = create_date_field(
        "Rotation End Date (YYYY-MM-DD)", 
        member_data.get("rotation_end_date"), 5
    )

    initial_active = bool(member_data.get("active", True))

    def on_active_toggle():
        current_active = fields["active"].get()
        if current_active == initial_active:
            return

        today_str = datetime.date.today().strftime("%Y-%m-%d")
        if current_active:
            fields["rotation_start_date"].delete(0, "end")
            fields["rotation_start_date"].insert(0, today_str)
            fields["rotation_end_date"].delete(0, "end")
        else:
            fields["rotation_end_date"].delete(0, "end")
            fields["rotation_end_date"].insert(0, today_str)
            fields["rotation_start_date"].delete(0, "end")

    fields["active"] = create_bool_field(
        "Active in Rotation",
        initial_active,
        3,
        command=on_active_toggle,
    )
    
    # Instructions
    instr = tk.Label(
        main_frame,
        text="Rotation order must be a positive integer. Dates are YYYY-MM-DD format.",
        font=("Arial", 9),
        fg="gray"
    )
    instr.grid(row=6, column=0, columnspan=2, sticky="w", pady=(8, 0))
    
    # Buttons
    button_frame = tk.Frame(main_frame)
    button_frame.grid(row=7, column=0, columnspan=2, pady=(12, 0))
    
    def save_changes():
        original_service_end = member_data.get("service_end_date") or ""
        new_service_end = (fields["service_end_date"].get() or "").strip()
        rotation_order = _coerce_positive_int((fields["rotation_order"].get() or "").strip())
        if rotation_order is None:
            show_focused_info_dialog(
                "Invalid rotation order",
                "Rotation Order # must be a positive integer.",
                dialog,
            )
            return

        updated = {
            "name": member_data["name"],
            "alias": member_data.get("alias") or member_data.get("initials"),
            "rotation_order": rotation_order,
            "service_start_date": fields["service_start_date"].get() or None,
            "service_end_date": new_service_end or None,
            "active": fields["active"].get(),
            "rotation_start_date": fields["rotation_start_date"].get() or None,
            "rotation_end_date": fields["rotation_end_date"].get() or None,
        }

        # If service end date changed to a non-empty value, force rotation to end on that date.
        if new_service_end and new_service_end != original_service_end:
            updated["active"] = False
            updated["rotation_start_date"] = None
            updated["rotation_end_date"] = new_service_end

        result["data"] = updated
        result["updated"] = True
        dialog.destroy()
    
    def cancel_edit():
        dialog.destroy()
    
    tk.Button(button_frame, text="Save", command=save_changes, width=10).pack(side="left", padx=4)
    tk.Button(button_frame, text="Cancel", command=cancel_edit, width=10).pack(side="left", padx=4)
    
    main_frame.columnconfigure(1, weight=1)
    
    dialog.protocol("WM_DELETE_WINDOW", cancel_edit)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()
    
    return result["data"] if result["updated"] else None


def show_member_management_dialog(parent=None):
    """
    Show a dialog to manage SWOC team members.
    Allows viewing, editing, adding members.
    """
    owner = parent or getattr(tk, "_default_root", None)
    dialog = tk.Toplevel(owner)
    dialog.title("SWOC Team Member Management")
    _set_dialog_geometry(dialog, 620, 480, min_width=520, min_height=380)
    dialog.resizable(True, True)
    
    members = load_swoc_members()
    
    # Capture initial state for change detection
    def get_member_list_state(member_list):
        return [(m.get("name"), m.get("rotation_order")) for m in member_list]
    
    initial_state = get_member_list_state(members)
    has_changes = [False]
    close_button_holder = {}

    def mark_changes():
        """Check for net changes and update button label."""
        current_state = get_member_list_state(members)
        has_changes[0] = current_state != initial_state
        if "widget" in close_button_holder and close_button_holder["widget"]:
            button_text = "Save" if has_changes[0] else "Continue"
            close_button_holder["widget"].config(text=button_text)

    def close_and_log():
        """Log member-order status and continue to SWOC schedule generation."""
        mark_changes()

        if has_changes[0]:
            save_swoc_members(members)

        log_file = write_swoc_member_list_log(members, changed=has_changes[0])
        _reveal_in_finder_and_refocus(log_file, dialog)

        dialog.destroy()
        run_swoc_schedule_generation(parent=None)

    
    # Create a frame with scrollbar
    main_frame = tk.Frame(dialog)
    main_frame.pack(fill="both", expand=True, padx=8, pady=8)
    
    # Header
    tk.Label(
        main_frame,
        text="Team Members",
        font=("Arial", 12, "bold")
    ).pack(anchor="w", pady=(0, 8))
    tk.Label(
        main_frame,
        text=(
            "Mouse: drag and drop to reorder.  Keyboard: arrows to select, "
            "Space to pick/drop, arrows to move while picked."
        ),
        font=("Arial", 9),
        fg="white",
    ).pack(anchor="w", pady=(0, 6))

    keyboard_status_var = tk.StringVar(value="Keyboard mode: use Up/Down to select a member.")
    tk.Label(
        main_frame,
        textvariable=keyboard_status_var,
        font=("Arial", 9),
        fg="white",
    ).pack(anchor="w", pady=(0, 6))
    
    # Listbox with members
    list_frame = tk.Frame(main_frame)
    list_frame.pack(fill="both", expand=True, pady=(0, 8))
    
    scrollbar = tk.Scrollbar(list_frame)
    scrollbar.pack(side="right", fill="y")
    
    listbox = tk.Listbox(list_frame, yscrollcommand=scrollbar.set, height=12)
    listbox.pack(side="left", fill="both", expand=True)
    scrollbar.config(command=listbox.yview)
    
    def refresh_listbox(selected_index=None):
        listbox.delete(0, "end")
        for i, member in enumerate(members):
            active_status = "✓" if member.get("active") else "✗"
            display_text = (
                f"#{i + 1:02d} "
                f"{member['name']} ({member.get('alias', '')}) [{active_status}]"
            )
            listbox.insert(i, display_text)
        if selected_index is not None and members:
            bounded_index = max(0, min(selected_index, len(members) - 1))
            listbox.selection_set(bounded_index)
            listbox.activate(bounded_index)

    refresh_listbox()
    if members:
        listbox.selection_set(0)
        listbox.activate(0)

    drag_start_index = [None]
    drag_current_index = [None]
    drag_start_order = [[]]
    drag_moved_name = [None]
    keyboard_reorder_active = [False]
    keyboard_start_order = [[]]
    keyboard_moved_name = [None]
    keyboard_pick_order = [[]]
    keyboard_pick_index = [0]

    def selected_index():
        selection = listbox.curselection()
        if not selection:
            if members:
                listbox.selection_set(0)
                listbox.activate(0)
                return 0
            return None
        return selection[0]

    def set_selected_index(index):
        if not members:
            return
        bounded_index = max(0, min(index, len(members) - 1))
        listbox.selection_clear(0, "end")
        listbox.selection_set(bounded_index)
        listbox.activate(bounded_index)
        listbox.see(bounded_index)

    def update_keyboard_status_for_selection(index):
        if index is None or not members:
            keyboard_status_var.set("Keyboard mode: no members available.")
            return
        member = members[index]
        if keyboard_reorder_active[0]:
            keyboard_status_var.set(
                f"Moving {member.get('name')}: use Up/Down to reposition, Space to drop."
            )
        else:
            keyboard_status_var.set(
                f"Selected #{index + 1:02d} {member.get('name')}. Press Space to pick up this member."
            )

    def commit_keyboard_reorder():
        members[:] = normalize_rotation_order(members, preserve_list_order=True)
        save_swoc_members(members)
        add_swoc_audit_log_entry(
            "MEMBER_REORDERED",
            keyboard_moved_name[0] or "UNKNOWN",
            str(keyboard_start_order[0] or []),
            str([m.get("name") for m in members]),
            "via keyboard reorder in member management interface",
        )
        mark_changes()

    def handle_up_down(delta):
        if not members:
            return "break"

        current = selected_index()
        if current is None:
            return "break"

        target = max(0, min(current + delta, len(members) - 1))
        if target == current:
            update_keyboard_status_for_selection(current)
            return "break"

        if keyboard_reorder_active[0]:
            moving_member = members.pop(current)
            members.insert(target, moving_member)
            refresh_listbox(selected_index=target)
            update_keyboard_status_for_selection(target)
        else:
            set_selected_index(target)
            update_keyboard_status_for_selection(target)

        return "break"

    def handle_space(_event):
        if not members:
            return "break"

        current = selected_index()
        if current is None:
            return "break"

        if not keyboard_reorder_active[0]:
            keyboard_reorder_active[0] = True
            keyboard_pick_order[0] = [dict(m) for m in members]
            keyboard_pick_index[0] = current
            keyboard_start_order[0] = [m.get("name") for m in members]
            keyboard_moved_name[0] = members[current].get("name")
            keyboard_status_var.set(
                f"Picked {members[current].get('name')}. Use Up/Down to move, Space to drop. Press ESC to cancel."
            )
            return "break"

        keyboard_reorder_active[0] = False
        commit_keyboard_reorder()
        refresh_listbox(selected_index=current)
        keyboard_status_var.set(
            f"Placed {keyboard_moved_name[0] or 'member'} at position #{current + 1:02d}."
        )
        keyboard_start_order[0] = []
        keyboard_moved_name[0] = None
        keyboard_pick_order[0] = []
        keyboard_pick_index[0] = 0
        return "break"

    def handle_escape(_event):
        if not keyboard_reorder_active[0]:
            return "break"

        if keyboard_pick_order[0]:
            members[:] = keyboard_pick_order[0]
            pick_index = keyboard_pick_index[0] or 0
            refresh_listbox(selected_index=pick_index)
            update_keyboard_status_for_selection(pick_index)

        keyboard_reorder_active[0] = False
        keyboard_pick_order[0] = []
        keyboard_pick_index[0] = 0
        keyboard_start_order[0] = []
        keyboard_moved_name[0] = None
        keyboard_status_var.set("Cancelled. Use Up/Down to select a member.")
        return "break"

    def begin_drag(event):
        keyboard_reorder_active[0] = False
        if not members:
            return
        index = listbox.nearest(event.y)
        if index < 0 or index >= len(members):
            return
        drag_start_index[0] = index
        drag_current_index[0] = index
        drag_start_order[0] = [m.get("name") for m in members]
        drag_moved_name[0] = members[index].get("name")
        listbox.selection_clear(0, "end")
        listbox.selection_set(index)
        update_keyboard_status_for_selection(index)

    def drag_member(event):
        start_index = drag_start_index[0]
        if start_index is None:
            return
        target_index = listbox.nearest(event.y)
        if target_index < 0 or target_index >= len(members):
            return
        current_index = drag_current_index[0]
        if current_index is None:
            return
        if target_index == current_index:
            return

        moving_member = members.pop(current_index)
        members.insert(target_index, moving_member)
        drag_current_index[0] = target_index
        refresh_listbox(selected_index=target_index)
        update_keyboard_status_for_selection(target_index)

    def end_drag(_event):
        start_index = drag_start_index[0]
        end_index = drag_current_index[0]
        if start_index is None or end_index is None:
            return

        if end_index != start_index:
            members[:] = normalize_rotation_order(members, preserve_list_order=True)
            save_swoc_members(members)
            add_swoc_audit_log_entry(
                "MEMBER_REORDERED",
                drag_moved_name[0] or "UNKNOWN",
                str(drag_start_order[0] or []),
                str([m.get("name") for m in members]),
                "via drag-and-drop in member management interface",
            )
            refresh_listbox(selected_index=end_index)
            update_keyboard_status_for_selection(end_index)
            mark_changes()

        drag_start_index[0] = None
        drag_current_index[0] = None
        drag_start_order[0] = []
        drag_moved_name[0] = None

    listbox.bind("<ButtonPress-1>", begin_drag)
    listbox.bind("<B1-Motion>", drag_member)
    listbox.bind("<ButtonRelease-1>", end_drag)
    listbox.bind("<Up>", lambda _event: handle_up_down(-1))
    listbox.bind("<Down>", lambda _event: handle_up_down(1))
    listbox.bind("<space>", handle_space)
    listbox.bind("<Escape>", handle_escape)
    update_keyboard_status_for_selection(selected_index())
    listbox.focus_set()

    # Buttons
    button_frame = tk.Frame(main_frame)
    button_frame.pack(fill="x", pady=(0, 8))

    def add_new_member():
        blank = {
            "name": "",
            "alias": "",
            "rotation_order": len(members) + 1,
            "service_start_date": datetime.date.today().isoformat(),
            "service_end_date": None,
            "active": True,
            "rotation_start_date": datetime.date.today().isoformat(),
            "rotation_end_date": None,
        }
        # Show a small dialog to collect name + alias first
        owner2 = dialog
        name_dialog = tk.Toplevel(owner2)
        name_dialog.title("New Member")
        _set_dialog_geometry(name_dialog, 380, 200, min_width=320, min_height=180)
        name_dialog.resizable(False, False)

        name_result = {"ok": False}
        nf = tk.Frame(name_dialog, padx=12, pady=12)
        nf.pack(fill="both", expand=True)

        tk.Label(nf, text="Full Name:").grid(row=0, column=0, sticky="w", pady=4)
        name_entry = tk.Entry(nf, width=24)
        name_entry.grid(row=0, column=1, sticky="ew", pady=4, padx=(8, 0))

        tk.Label(nf, text="Alias:").grid(row=1, column=0, sticky="w", pady=4)
        alias_entry = tk.Entry(nf, width=12)
        alias_entry.grid(row=1, column=1, sticky="w", pady=4, padx=(8, 0))

        bf2 = tk.Frame(nf)
        bf2.grid(row=2, column=0, columnspan=2, pady=(10, 0))

        def confirm_name():
            n = name_entry.get().strip()
            alias = alias_entry.get().strip()
            if not n or not alias:
                show_focused_info_dialog("Required", "Name and alias are required.", name_dialog)
                return
            blank["name"] = n
            blank["alias"] = alias
            name_result["ok"] = True
            name_dialog.destroy()

        def cancel_name():
            name_dialog.destroy()

        tk.Button(bf2, text="Next", command=confirm_name, width=8).pack(side="left", padx=4)
        tk.Button(bf2, text="Cancel", command=cancel_name, width=8).pack(side="left", padx=4)
        nf.columnconfigure(1, weight=1)
        name_dialog.protocol("WM_DELETE_WINDOW", cancel_name)
        name_dialog.lift()
        name_dialog.focus_force()
        name_dialog.attributes("-topmost", True)
        name_dialog.grab_set()
        name_dialog.wait_window()

        if not name_result["ok"]:
            return

        updated = show_member_edit_dialog(blank["name"], blank, dialog)
        if updated:
            # New members are always appended to the end of the current rotation.
            updated["rotation_order"] = len(members) + 1
            members.append(updated)
            members[:] = normalize_rotation_order(members, preserve_list_order=True)
            save_swoc_members(members)
            add_swoc_audit_log_entry(
                "MEMBER_ADDED",
                updated["name"],
                None,
                str(updated),
                "via member management interface",
            )
            refresh_listbox()
            mark_changes()
            show_focused_info_dialog("Success", f"{updated['name']} added and logged to audit trail.", dialog)

    def edit_member():
        selection = listbox.curselection()
        if not selection:
            show_focused_info_dialog("No Selection", "Please select a member to edit.", dialog)
            return
        
        idx = selection[0]
        member_to_edit = members[idx]
        
        updated = show_member_edit_dialog(
            member_to_edit["name"],
            member_to_edit,
            dialog
        )
        
        if updated:
            old_member = dict(members[idx])
            members[idx] = updated
            requested_order = _coerce_positive_int(updated.get("rotation_order")) or (idx + 1)
            requested_order = max(1, min(requested_order, len(members)))

            if requested_order != idx + 1:
                moved_member = members.pop(idx)
                members.insert(requested_order - 1, moved_member)

            members[:] = normalize_rotation_order(members, preserve_list_order=True)
            save_swoc_members(members)
            add_swoc_audit_log_entry(
                "MEMBER_UPDATED",
                updated["name"],
                str(old_member),
                str(updated),
                "via member management interface",
            )
            refresh_listbox()
            mark_changes()
            show_focused_info_dialog("Success", "Member updated and logged to audit trail.", dialog)

    def remove_member():
        selection = listbox.curselection()
        if not selection:
            show_focused_info_dialog("No Selection", "Please select a member to remove.", dialog)
            return

        idx = selection[0]
        member_to_remove = members[idx]
        if not show_focused_yes_no_dialog(
            "Remove Member",
            (
                f"Remove {member_to_remove['name']} from rotation?\n\n"
                "This will re-number all following members."
            ),
            parent=dialog,
        ):
            return

        removed_member = members.pop(idx)
        members[:] = normalize_rotation_order(members, preserve_list_order=True)
        save_swoc_members(members)
        add_swoc_audit_log_entry(
            "MEMBER_REMOVED",
            removed_member["name"],
            str(removed_member),
            None,
            "via member management interface",
        )
        refresh_listbox()
        mark_changes()
        show_focused_info_dialog(
            "Success",
            f"{removed_member['name']} removed; rotation order was re-numbered.",
            dialog,
        )

    tk.Button(button_frame, text="Add New", command=add_new_member).pack(side="left", padx=4)
    tk.Button(button_frame, text="Edit Selected", command=edit_member).pack(side="left", padx=4)
    tk.Button(button_frame, text="Remove Selected", command=remove_member).pack(side="left", padx=4)
    close_button_holder["widget"] = tk.Button(button_frame, text="Continue", command=close_and_log)
    close_button_holder["widget"].pack(side="right", padx=4)
    
    dialog.protocol("WM_DELETE_WINDOW", close_and_log)
    dialog.lift()
    dialog.focus_force()
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.wait_window()


def main():
    root = tk.Tk()
    root.title("Telescope & Staff Schedules")
    root.geometry("640x420")
    root.resizable(False, False)
    startup_action = tk.StringVar(value="")
    mode_var = tk.StringVar(value="excel")

    tk.Label(
        root,
        text="Select a Schedule Type",
        padx=16,
        pady=12,
        font=("Arial", 12, "bold"),
    ).pack(anchor="w")

    tk.Label(
        root,
        text="Step 1: Telescope Schedule",
        padx=16,
        pady=4,
        font=("Arial", 11, "bold"),
    ).pack(anchor="w")

    telescope_radio_frame = tk.Frame(root, padx=24)
    telescope_radio_frame.pack(anchor="w", pady=(0, 10))
    tk.Radiobutton(
        telescope_radio_frame,
        text="Telescope Schedule - Import from Excel file",
        variable=mode_var,
        value="telescope",
    ).pack(anchor="w", pady=2)

    tk.Label(
        root,
        text="Step 11: Night Staff Schedule:",
        padx=16,
        pady=4,
        font=("Arial", 11, "bold"),
    ).pack(anchor="w")

    radio_frame = tk.Frame(root, padx=24)
    radio_frame.pack(anchor="w", pady=4)
    tk.Radiobutton(
        radio_frame, text="SA / OA / NA — Import from Excel file", variable=mode_var, value="excel"
    ).pack(anchor="w", pady=2)
    tk.Radiobutton(
        radio_frame, text="SWOC — Generate rotation schedule", variable=mode_var, value="swoc"
    ).pack(anchor="w", pady=2)

    tk.Label(
        root,
        text=(
            "The telescope workflow validates a telescope schedule and generates SQL INSERT statements.\n"
            "The staff workflows generate SQL INSERT statements and update the Night Staff Schedule database."
        ),
        padx=16,
        pady=8,
        justify="left",
        fg="gray",
    ).pack(anchor="w")

    def select_mode():
        startup_action.set(mode_var.get())

    def cancel_startup():
        startup_action.set("cancel")

    button_frame = tk.Frame(root)
    button_frame.pack(pady=10)
    tk.Button(button_frame, text="OK", command=select_mode).pack(side="left", padx=6)
    tk.Button(button_frame, text="Cancel", command=cancel_startup).pack(side="left", padx=6)

    root.protocol("WM_DELETE_WINDOW", cancel_startup)

    root.lift()
    root.focus_force()
    root.attributes("-topmost", True)
    root.update()

    root.wait_variable(startup_action)
    chosen_mode = startup_action.get()
    if chosen_mode == "cancel":
        root.destroy()
        return

    if chosen_mode == "telescope":
        root.destroy()
        from s01_tel_sched import step01_tel_sched

        step01_tel_sched.main()
        return

    # ── SWOC mode ──────────────────────────────────────────────────────────
    if chosen_mode == "swoc":
        # Iconify instead of withdraw so the dialog has a live parent on macOS
        root.iconify()
        show_member_management_dialog(parent=root)
        root.destroy()
        return

    # ── Excel import mode (SA / OA / NA) ───────────────────────────────────
    upload_enabled = True
    root.withdraw()
    file_path = pick_file(root)
    if not file_path:
        root.destroy()
        return

    # Detect and confirm staff type
    detected_staff_type = _normalize_staff_type(Path(file_path).stem)
    if not confirm_staff_type_dialog(detected_staff_type or "unknown", file_path, parent=root):
        root.destroy()
        return

    data_dir = ensure_data_dir()
    copied_file_path = copy_source_to_data_dir(file_path, data_dir)

    sheets = get_sheet_names(copied_file_path)

    # Open spreadsheet visually
    calc_process = open_in_calc(copied_file_path, sheets[0])

    # Let user choose
    chosen_sheet = choose_sheet(
        sheets,
        file_path=copied_file_path,
        default_sheet=sheets[0],
        parent=root,
    )
    if chosen_sheet is None:
        root.destroy()
        return

    # Export selected sheet
    result = export_sheet(copied_file_path, chosen_sheet, data_dir, upload_enabled, staff_type=detected_staff_type)
    if not result:
        root.destroy()
        return
    _, output, upload_message, sql_file = result

    completion_action = show_final_completion_dialog(
        output=output,
        upload_message=upload_message,
        sql_file=sql_file,
        parent=root,
    )
    close_calc_process(calc_process)

    if completion_action != "ok":
        root.destroy()
        return

    if not preflight_mysql_connection_check():
        root.destroy()
        return

    if not connect_to_remote_mysql_db():
        root.destroy()
        return

    if sql_file:
        focus_converted_file(sql_file)
    root.destroy()

if __name__ == "__main__":
    main()